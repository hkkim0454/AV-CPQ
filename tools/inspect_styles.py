# -*- coding: utf-8 -*-
"""원본의 셀 서식(폰트/테두리/정렬/표시형식)만 추출한다. 값은 출력하지 않는다."""
import sys
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8")

def sig(c):
    f, a, b, fill = c.font, c.alignment, c.border, c.fill
    def side(s): return f"{s.style or '-'}" 
    return (
        f"font({f.name},{f.sz},{'B' if f.b else ''}{'I' if f.i else ''},{f.color.rgb if f.color and f.color.type=='rgb' else '-'})"
        f" align({a.horizontal or '-'},{a.vertical or '-'},wrap={bool(a.wrapText)},shrink={bool(a.shrinkToFit)},ind={a.indent})"
        f" border(l={side(b.left)},r={side(b.right)},t={side(b.top)},b={side(b.bottom)})"
        f" fill({fill.fgColor.rgb if fill and fill.patternType else '-'})"
        f" fmt({c.number_format})"
    )

wb = load_workbook(sys.argv[1])
for ws in wb.worksheets:
    print(f"\n===== {ws.title!r} =====")
    seen = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), max_col=11):
        for c in row:
            s = sig(c)
            seen.setdefault(s, []).append(c.coordinate)
    for s, coords in seen.items():
        print(f"{s}\n    -> {len(coords)} cells: {','.join(coords[:14])}{' ...' if len(coords)>14 else ''}")
