# -*- coding: utf-8 -*-
"""`tools/extract_catalog.py` 의 시험 (계획 2026-10-05 품셈 교체 Task 1).

**합성 시트만 쓴다.** 실제 품셈 원본을 열지 않는다 — 시험이 사용자 PC의
파일에 기대면 다른 곳에서 돌아가지 않고, 실패 메시지에 원본 값이 섞인다.

pytest 를 쓰지 않는 이유: 이 저장소에는 파이썬 시험 묶음이 없다.
의존성을 하나 더 들이는 대신 이 파일 하나로 돌린다.

    python tools/test_extract_catalog.py
"""
import json
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter as gl

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.stdout.reconfigure(encoding="utf-8")

import extract_catalog as ec


# --- 아주 작은 시험 틀 -------------------------------------------------------
_FAILURES: list[str] = []
_PASSED = 0


def check(name, fn):
    global _PASSED
    try:
        fn()
    except AssertionError as error:
        _FAILURES.append(f"{name}\n      {error}")
        print(f"  x  {name}")
    except Exception as error:  # noqa: BLE001 — 시험 틀이므로 전부 받는다
        _FAILURES.append(f"{name}\n      예상 못한 예외: {type(error).__name__}: {error}")
        print(f"  x  {name}  ({type(error).__name__})")
    else:
        _PASSED += 1
        print(f"  ok {name}")


def expect_system_exit(fn):
    """`SystemExit` 를 받아 메시지를 돌려준다. 안 나면 단언 실패."""
    try:
        fn()
    except SystemExit as exit_error:
        return str(exit_error)
    raise AssertionError("SystemExit 로 중단해야 하는데 그냥 끝났다")


# --- 합성 통합문서 -----------------------------------------------------------
#
# 두 판의 머리글 배치를 **실측 그대로** 흉내낸다. 열 번호가 서로 다르다는
# 것이 요점이다 — 머리글로 찾지 않으면 한쪽에서만 맞는다.
#
#   상반기  G 재료비/단가   I 노무비/단가   M 제조사  N 영업비고  P 코드  Q 요율
#   하반기  G 원가/단가     I 재료비/단가   K 노무비/단가  P 제조사  Q 영업비고  S 코드  T 요율

H2_LAYOUT = {
    "row2": {1: "번호", 2: "품   명", 3: "규 격", 4: "설 명", 5: "단위", 6: "수량",
             7: "원가", 9: "재료비", 11: "노무비", 13: "합 계", 14: "재료비 이윤",
             15: "비 고", 16: "제조사/구매처", 17: "영업비고", 18: "노무비",
             20: "품목별 요율%", 21: "할증", 22: "26년 하반기"},
    "row3": {7: "단 가", 8: "금액", 9: "단 가", 10: "금액", 11: "단 가", 12: "금액",
             21: 0, 22: "표준단가"},
    "merges": ["G2:H2", "I2:J2", "K2:L2", "A2:A3", "B2:B3", "C2:C3", "D2:D3", "E2:E3"],
    "trades": {23: ("통신내선공", "M/D", 286219), 25: ("통신설비공", "M/D", 300000)},
    "cols": {"selling": 9, "labor": 11, "code": 19, "rate": 20, "surcharge": 21,
             "maker": 16, "sales": 17, "remark": 15, "trade1": 23},
}

H1_LAYOUT = {
    "row2": {1: "번호", 2: "품   명", 3: "규 격", 4: "설 명", 5: "단위", 6: "수량",
             7: "재료비", 9: "노무비", 11: "합 계", 12: "비 고",
             13: "제조사/구매처", 14: "영업비고", 15: "노무비",
             17: "품목별 요율%", 18: "할증", 19: "26년 상반기"},
    "row3": {7: "단 가", 8: "금액", 9: "단 가", 10: "금액", 18: 0, 19: "표준단가"},
    "merges": ["G2:H2", "I2:J2", "A2:A3", "B2:B3", "C2:C3", "D2:D3", "E2:E3"],
    "trades": {20: ("통신관련 기사", "M/D", 320449), 22: ("통신설비공", "M/D", 300000)},
    "cols": {"selling": 7, "labor": 9, "code": 16, "rate": 17, "surcharge": 18,
             "maker": 13, "sales": 14, "remark": 12, "trade1": 20},
}


def write_sheet(ws, layout, body=()):
    for col, value in layout["row2"].items():
        ws.cell(row=2, column=col).value = value
    for col, value in layout["row3"].items():
        ws.cell(row=3, column=col).value = value
    for rng in layout["merges"]:
        ws.merge_cells(rng)
    for col, (name, unit, wage) in layout["trades"].items():
        ws.cell(row=2, column=col).value = name
        ws.cell(row=3, column=col).value = unit
        ws.cell(row=3, column=col + 1).value = wage
    for offset, product in enumerate(body):
        write_product(ws, 4 + offset, layout, **product)
    return ws


def write_product(ws, row, layout, *, name, spec="SPEC", unit="EA", selling=None,
                  labor=None, code=None, rate=None, surcharge=None, trade1=None,
                  maker=None, sales_note=None, description=None, remark=None,
                  trade1_amount=None, trade2=None, trade2_amount=None):
    """제품 한 행. `maker`/`sales_note` 는 **금지 열을 일부러 채우는** 시험용이다."""
    c = layout["cols"]
    ws.cell(row=row, column=1).value = row
    ws.cell(row=row, column=2).value = name
    ws.cell(row=row, column=3).value = spec
    ws.cell(row=row, column=4).value = description
    ws.cell(row=row, column=5).value = unit
    ws.cell(row=row, column=c["selling"]).value = selling
    ws.cell(row=row, column=c["labor"]).value = labor
    ws.cell(row=row, column=c["remark"]).value = remark
    ws.cell(row=row, column=c["maker"]).value = maker
    ws.cell(row=row, column=c["sales"]).value = sales_note
    ws.cell(row=row, column=c["code"]).value = code
    ws.cell(row=row, column=c["rate"]).value = rate
    ws.cell(row=row, column=c["surcharge"]).value = surcharge
    for offset, (quantity, amount) in enumerate(((trade1, trade1_amount), (trade2, trade2_amount))):
        if quantity is None:
            continue
        q_col = c["trade1"] + offset * 2
        a_col = q_col + 1
        ws.cell(row=row, column=q_col).value = quantity
        # 직종별 '금액' 칸. 원본의 표준은 `=공수*노임` 이고, 일부 행은
        # `=INT(공수*노임)` 이며, 상수로 덮인 행도 있다.
        plain = f"={gl(q_col)}{row}*{gl(a_col)}$3"
        ws.cell(row=row, column=a_col).value = (
            plain if amount is None
            else f"=INT({gl(q_col)}{row}*{gl(a_col)}$3)" if amount == "int"
            else plain if amount == "plain"
            else f"={gl(q_col)}{row}*오디오!{gl(a_col)}$3" if amount == "cross-sheet"
            else amount
        )


SAMPLE = dict(name="합성 스위처", spec="SYN-100", selling=1000000, labor=50000,
              code="7-11-1-합성 설치", rate=1, surcharge=0, trade1=0.5,
              maker="합성제조사/합성구매처", sales_note="260101_단가확인 합성현장",
              remark="합성 비고")

LEAK_MARKERS = ("합성제조사", "합성구매처", "합성현장", "260101")


def build(tmp, layout=H2_LAYOUT, body=(SAMPLE,), title="영상", extra=None):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    write_sheet(ws, layout, body)
    if extra is not None:
        extra(ws, wb)
    path = Path(tmp) / "synthetic.xlsx"
    wb.save(path)
    wb.close()
    return path


def extract_to_dict(tmp, **kwargs):
    source = build(tmp, **kwargs)
    destination = Path(tmp) / "out.json"
    ec.extract(source, destination)
    return json.loads(destination.read_text(encoding="utf-8")), destination


# --- 시험: 올바른 열을 읽는다 -------------------------------------------------
def _assert_reads_correctly(layout):
    with tempfile.TemporaryDirectory() as tmp:
        data, destination = extract_to_dict(tmp, layout=layout)
        row = data["sheets"][0]["rows"][0]
        assert row["materialUnitPrice"] == "1000000", f"판매단가: {row.get('materialUnitPrice')!r}"
        assert row["laborUnitPrice"] == "50000", f"노무비 단가: {row.get('laborUnitPrice')!r}"
        assert row["laborCode"] == "7-11-1-합성 설치", f"품셈 코드: {row.get('laborCode')!r}"
        assert row["itemRate"] == "1", f"품목별 요율: {row.get('itemRate')!r}"
        assert row["surcharge"] == "0", f"할증: {row.get('surcharge')!r}"
        assert row["remark"] == "합성 비고", f"비고: {row.get('remark')!r}"
        assert row["trades"] == [{"trade": "통신내선공" if layout is H2_LAYOUT else "통신관련 기사",
                                  "quantity": "0.5"}], f"직종: {row.get('trades')!r}"
        text = destination.read_text(encoding="utf-8")
        for marker in LEAK_MARKERS:
            assert marker not in text, f"금지 열의 값이 결과에 들어갔다: {marker!r}"


def test_reads_h2_layout():
    """하반기 배치에서 머리글로 올바른 열을 읽는다 (예전 코드는 16열을 품셈 코드로 읽었다)."""
    _assert_reads_correctly(H2_LAYOUT)


def test_reads_h1_layout():
    """상반기 배치에서도 같은 코드가 동작한다 — 열 번호를 외우지 않는다는 증거."""
    _assert_reads_correctly(H1_LAYOUT)


def test_wages_use_column_plus_one():
    """직종 열의 `col+1` 이 노임이라는 좌표 계약."""
    with tempfile.TemporaryDirectory() as tmp:
        data, _ = extract_to_dict(tmp)
        assert data["sheets"][0]["wages"] == [
            {"trade": "통신내선공", "unit": "M/D", "amount": "286219", "quantityColumn": "W"},
            {"trade": "통신설비공", "unit": "M/D", "amount": "300000", "quantityColumn": "Y"},
        ], data["sheets"][0]["wages"]


def test_labor_code_shape_is_not_restricted():
    """품셈 코드를 `7-11-1` 꼴로 한정하지 않는다. 덤프는 판단하지 않고 그대로 옮긴다."""
    codes = ["9-2-1-1 촬상부-브라켓-일반형", "8-1-1-Slot Type 장비설치_Card설치(Module)", "7-11-1"]
    body = [dict(SAMPLE, name=f"합성 {i}", code=code) for i, code in enumerate(codes)]
    with tempfile.TemporaryDirectory() as tmp:
        data, _ = extract_to_dict(tmp, body=body)
        assert [r["laborCode"] for r in data["sheets"][0]["rows"]] == codes


# --- 시험: 금지 머리글 -------------------------------------------------------
def _forbidden_layout(slot):
    """읽기 대상 열에 금지 머리글을 밀어 넣은 배치를 만든다."""
    layout = {k: (dict(v) if isinstance(v, dict) else list(v)) for k, v in H2_LAYOUT.items()}
    layout["cols"] = dict(H2_LAYOUT["cols"])
    layout["trades"] = dict(H2_LAYOUT["trades"])
    if slot == "trade":
        # 첫 직종 이름 자리에 금지 머리글이 온다.
        layout["trades"] = {23: ("제조사/구매처", "M/D", 286219), 25: ("통신설비공", "M/D", 300000)}
    elif slot == "code":
        # 머리글이 없어야 할 품셈 코드 자리에 금지 머리글이 온다.
        layout["row2"] = dict(H2_LAYOUT["row2"])
        layout["row2"][19] = "영업비고"
    return layout


def test_forbidden_header_in_trade_slot_aborts():
    """금지 머리글이 **직종 이름 자리**(읽기 대상)에 오면 전체 중단."""
    with tempfile.TemporaryDirectory() as tmp:
        destination = Path(tmp) / "out.json"
        source = build(tmp, layout=_forbidden_layout("trade"))
        message = expect_system_exit(lambda: ec.extract(source, destination))
        assert "제조사" in message or "구매처" in message, message
        assert not destination.exists(), "중단했으면 결과 파일을 쓰지 않아야 한다"


def test_forbidden_header_in_code_slot_aborts():
    """금지 머리글이 **품셈 코드 자리**(읽기 대상)에 오면 전체 중단."""
    with tempfile.TemporaryDirectory() as tmp:
        destination = Path(tmp) / "out.json"
        source = build(tmp, layout=_forbidden_layout("code"))
        message = expect_system_exit(lambda: ec.extract(source, destination))
        assert "영업비고" in message, message
        assert not destination.exists()


def test_abort_message_has_no_cell_values():
    """중단 메시지에 **그 열의 값**을 적지 않는다. 열 문자와 머리글만 적는다."""
    for slot in ("trade", "code"):
        with tempfile.TemporaryDirectory() as tmp:
            source = build(tmp, layout=_forbidden_layout(slot))
            message = expect_system_exit(lambda: ec.extract(source, Path(tmp) / "out.json"))
            for marker in LEAK_MARKERS:
                assert marker not in message, f"[{slot}] 메시지에 셀 값이 샜다: {marker!r} / {message!r}"


# --- 시험: 결손 · 모호 · 구조 ------------------------------------------------
def test_missing_header_aborts():
    """필수 머리글이 없으면 중단한다. 그 열만 건너뛰지 않는다."""
    layout = {k: (dict(v) if isinstance(v, dict) else list(v)) for k, v in H2_LAYOUT.items()}
    layout["row2"] = dict(H2_LAYOUT["row2"])
    del layout["row2"][9]   # 재료비(판매단가) 머리글을 없앤다
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, layout=layout)
        message = expect_system_exit(lambda: ec.extract(source, Path(tmp) / "out.json"))
        assert "재료비" in message, message


def test_duplicate_header_aborts():
    """같은 머리글이 두 열에 있으면 어느 쪽인지 정하지 않고 중단한다."""
    layout = {k: (dict(v) if isinstance(v, dict) else list(v)) for k, v in H2_LAYOUT.items()}
    layout["row2"] = dict(H2_LAYOUT["row2"])
    layout["row3"] = dict(H2_LAYOUT["row3"])
    layout["row2"][13] = "재료비"      # 합 계 자리를 재료비로 바꾸고
    layout["row3"][13] = "단 가"       # 하위도 단가로 맞춰 둘이 겹치게 한다
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, layout=layout)
        message = expect_system_exit(lambda: ec.extract(source, Path(tmp) / "out.json"))
        assert "여러 열" in message, message
        assert "I" in message and "M" in message, message


def test_labor_block_must_be_consecutive():
    """[코드][요율][할증][표준단가] 가 이 순서로 연속하지 않으면 중단한다."""
    layout = {k: (dict(v) if isinstance(v, dict) else list(v)) for k, v in H2_LAYOUT.items()}
    layout["row2"] = dict(H2_LAYOUT["row2"])
    del layout["row2"][21]            # 할증을 요율 바로 오른쪽에서 치우고
    layout["row2"][13] = "할증"        # 엉뚱한 자리에 둔다
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, layout=layout)
        message = expect_system_exit(lambda: ec.extract(source, Path(tmp) / "out.json"))
        assert "할증" in message and "오른쪽" in message, message


def test_header_in_code_slot_aborts():
    """품셈 코드 자리에 (금지어가 아니어도) 머리글이 있으면 구조가 다른 것이다."""
    layout = {k: (dict(v) if isinstance(v, dict) else list(v)) for k, v in H2_LAYOUT.items()}
    layout["row2"] = dict(H2_LAYOUT["row2"])
    layout["row2"][19] = "참고"
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, layout=layout)
        message = expect_system_exit(lambda: ec.extract(source, Path(tmp) / "out.json"))
        assert "품셈 코드 자리" in message, message


def test_missing_wage_aborts():
    """직종은 있는데 `col+1` 에 노임이 없으면 중단한다 — 노임이 비면 노무비가 0이 된다."""
    def blank_wage(ws, wb):
        ws.cell(row=3, column=24).value = None
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, extra=blank_wage)
        message = expect_system_exit(lambda: ec.extract(source, Path(tmp) / "out.json"))
        assert "노임" in message, message


# --- 시험: 정규화 · 병합 -----------------------------------------------------
def test_header_normalization():
    """공백·개행·전각을 정규화한 뒤 비교한다."""
    layout = {k: (dict(v) if isinstance(v, dict) else list(v)) for k, v in H2_LAYOUT.items()}
    layout["row2"] = dict(H2_LAYOUT["row2"])
    layout["row3"] = dict(H2_LAYOUT["row3"])
    layout["row2"][9] = "재 료\n비"        # 개행 + 공백
    layout["row3"][9] = "단　가"            # 전각 공백
    layout["row2"][20] = "품목별 요율％"     # 전각 퍼센트
    with tempfile.TemporaryDirectory() as tmp:
        data, _ = extract_to_dict(tmp, layout=layout)
        assert data["sheets"][0]["rows"][0]["materialUnitPrice"] == "1000000"


def test_merge_does_not_forward_fill():
    """병합은 anchor 범위 안에서만 펼친다. 범위 밖 열이 이름을 얻으면 안 된다.

    `노무비`(K2:L2) 아래 L3 는 `금액`이다. forward fill 을 하면 M·N 열까지
    `노무비`가 번져 `노무비+단가` 판정이 모호해진다.
    """
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp)
        from openpyxl import load_workbook
        ws = load_workbook(source).active
        roles = ec.header_grid(ws, 2, 3)
        assert roles[11] == ("노무비", "단가"), roles[11]
        assert roles[12] == ("노무비", "금액"), roles[12]
        assert roles[13][0] == "합계", roles[13]      # 병합 범위 밖 — 번지지 않았다
        assert roles[2] == ("품명", ""), roles[2]      # 세로 병합은 하위를 비운다


def test_finds_header_row():
    """머리글 행을 찾는다. `2·3행`이라고 외우지 않는다."""
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp)
        from openpyxl import load_workbook
        ws = load_workbook(source).active
        assert ec.find_header_rows(ws, "영상") == (2, 3, 4)


# --- 시험: LED 시트 거르기 ---------------------------------------------------
def test_led_sheet_skipped_even_when_visible():
    """`LED전광판 계산` 은 **이름으로** 거른다. 하반기 파일에서 visible 로 바뀌었다."""
    def add_led(ws, wb):
        led = wb.create_sheet("LED전광판 계산")
        write_sheet(led, H2_LAYOUT, [dict(SAMPLE, name="합성 LED 캐비넷")])
        assert led.sheet_state == "visible"
    with tempfile.TemporaryDirectory() as tmp:
        data, destination = extract_to_dict(tmp, extra=add_led)
        names = [s["name"] for s in data["sheets"]]
        assert names == ["영상"], names
        assert "합성 LED 캐비넷" not in destination.read_text(encoding="utf-8")


# --- 시험: 노무비 단가 수식의 배율 -------------------------------------------
#
# 원본 일부 행은 `=INT(SUM((할증*표준단가),표준단가)*요율)*0.3` 처럼 **열에 없는
# 배율**이 수식 끝에 붙어 있다(실측 18행). 그 배율을 읽지 못하면 노무비가
# 최대 3.3배 부풀어 들어간다.
def labor_formula(row, layout, multiplier=None):
    c = layout["cols"]
    base = (f"=INT(SUM(({gl(c['surcharge'])}{row}*{gl(c['surcharge'] + 1)}{row}),"
            f"{gl(c['surcharge'] + 1)}{row})*{gl(c['rate'])}{row})")
    return base if multiplier is None else f"{base}*{multiplier}"


def test_reads_multiplier_from_formula():
    """수식 끝의 배율을 읽어 덤프에 담는다."""
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, body=[dict(SAMPLE, labor=None)])
        from openpyxl import load_workbook
        wb = load_workbook(source)
        ws = wb.active
        ws.cell(row=4, column=H2_LAYOUT["cols"]["labor"]).value = labor_formula(4, H2_LAYOUT, "0.3")
        wb.save(source)
        wb.close()
        destination = Path(tmp) / "out.json"
        ec.extract(source, destination)
        rec = json.loads(destination.read_text(encoding="utf-8"))["sheets"][0]["rows"][0]
        assert rec.get("laborMultiplier") == "0.3", f"배율을 못 읽었다: {rec.get('laborMultiplier')!r}"
        assert "laborFormulaUnrecognized" not in rec, rec


def test_no_multiplier_on_plain_formula():
    """배율이 없는 표준 수식에는 배율을 붙이지 않는다."""
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, body=[dict(SAMPLE, labor=None)])
        from openpyxl import load_workbook
        wb = load_workbook(source)
        wb.active.cell(row=4, column=H2_LAYOUT["cols"]["labor"]).value = labor_formula(4, H2_LAYOUT)
        wb.save(source)
        wb.close()
        destination = Path(tmp) / "out.json"
        ec.extract(source, destination)
        rec = json.loads(destination.read_text(encoding="utf-8"))["sheets"][0]["rows"][0]
        assert "laborMultiplier" not in rec, rec
        assert "laborFormulaUnrecognized" not in rec, rec


def test_unknown_formula_is_flagged_not_guessed():
    """모르는 모양이면 **해석하지 않고 표시만** 한다. 일반 계산기를 만들지 않는다."""
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, body=[dict(SAMPLE, labor=None)])
        from openpyxl import load_workbook
        wb = load_workbook(source)
        wb.active.cell(row=4, column=H2_LAYOUT["cols"]["labor"]).value = "=ROUNDUP(G4*1.1,-3)"
        wb.save(source)
        wb.close()
        destination = Path(tmp) / "out.json"
        ec.extract(source, destination)
        rec = json.loads(destination.read_text(encoding="utf-8"))["sheets"][0]["rows"][0]
        assert rec.get("laborFormulaUnrecognized") is True, rec
        assert "laborMultiplier" not in rec, rec


def test_multiplier_is_not_guessed_from_spec():
    """⛔ 규격 문자열(`SM 4C-30m`)에서 배율을 유추하지 않는다."""
    with tempfile.TemporaryDirectory() as tmp:
        source = build(tmp, body=[dict(SAMPLE, spec="SM 4C-30m", labor=None)])
        from openpyxl import load_workbook
        wb = load_workbook(source)
        wb.active.cell(row=4, column=H2_LAYOUT["cols"]["labor"]).value = labor_formula(4, H2_LAYOUT)
        wb.save(source)
        wb.close()
        destination = Path(tmp) / "out.json"
        ec.extract(source, destination)
        rec = json.loads(destination.read_text(encoding="utf-8"))["sheets"][0]["rows"][0]
        assert "laborMultiplier" not in rec, "규격에서 0.3 을 유추하면 안 된다"


# --- 시험: 직종 금액 칸의 모양 ------------------------------------------------
#
# 실측(하반기 전수): `=공수*노임` 3,239칸 · `=공수*다른시트!노임` 21칸 ·
# `=INT(공수*노임)` 3칸(오디오 371행 한 행) · 상수 2칸(CMS 9행 한 행).
#
# ⛔ **일반 규칙으로 만들지 않는다.** 그 행의 칸이 실제로 그 모양일 때만
# 그 행에 적용한다. 섞여 있으면 아는 모양이 아니므로 해석하지 않는다.
def shape_of(tmp, **product):
    data, _ = extract_to_dict(tmp, body=[dict(SAMPLE, **product)])
    return data["sheets"][0]["rows"][0].get("tradeAmountShape")


def test_plain_amounts_have_no_shape():
    """표준(`=공수*노임`)은 아무 표시도 하지 않는다."""
    with tempfile.TemporaryDirectory() as tmp:
        assert shape_of(tmp, trade1=0.5, trade2=0.2) is None


def test_cross_sheet_amounts_are_still_plain():
    """`=공수*다른시트!노임` 도 표준이다 — 노임 값은 전 시트가 같다."""
    with tempfile.TemporaryDirectory() as tmp:
        assert shape_of(tmp, trade1=0.5, trade1_amount="cross-sheet") is None


def test_all_int_amounts_are_marked_int():
    """그 행의 금액 칸이 **전부** `=INT(공수*노임)` 이면 `int` 로 표시한다."""
    with tempfile.TemporaryDirectory() as tmp:
        assert shape_of(tmp, trade1=0.5, trade1_amount="int",
                        trade2=0.2, trade2_amount="int") == "int"


def test_mixed_amounts_are_not_int():
    """⛔ 일부만 `INT` 인 행은 **아는 모양이 아니다.** 섞인 채로 해석하지 않는다."""
    with tempfile.TemporaryDirectory() as tmp:
        assert shape_of(tmp, trade1=0.5, trade1_amount="int",
                        trade2=0.2, trade2_amount="plain") == "mixed"


def test_constant_amount_is_marked_constant():
    """상수로 덮인 칸이 있으면 `constant` 다 (실측: CMS 9행)."""
    with tempfile.TemporaryDirectory() as tmp:
        assert shape_of(tmp, trade1=0.5, trade1_amount=31618.3) == "constant"


def main():
    print("tools/test_extract_catalog.py")
    check("하반기 배치에서 머리글로 올바른 열을 읽는다", test_reads_h2_layout)
    check("상반기 배치에서도 같은 코드가 동작한다", test_reads_h1_layout)
    check("직종 열의 col+1 이 노임이라는 좌표 계약", test_wages_use_column_plus_one)
    check("품셈 코드 모양을 한정하지 않는다", test_labor_code_shape_is_not_restricted)
    check("금지 머리글이 직종 이름 자리에 오면 중단한다", test_forbidden_header_in_trade_slot_aborts)
    check("금지 머리글이 품셈 코드 자리에 오면 중단한다", test_forbidden_header_in_code_slot_aborts)
    check("중단 메시지에 셀 값을 적지 않는다", test_abort_message_has_no_cell_values)
    check("필수 머리글이 없으면 중단한다", test_missing_header_aborts)
    check("같은 머리글이 두 열에 있으면 중단한다", test_duplicate_header_aborts)
    check("품셈 블록이 연속하지 않으면 중단한다", test_labor_block_must_be_consecutive)
    check("품셈 코드 자리에 머리글이 있으면 중단한다", test_header_in_code_slot_aborts)
    check("직종의 노임이 없으면 중단한다", test_missing_wage_aborts)
    check("공백·개행·전각을 정규화해 찾는다", test_header_normalization)
    check("병합을 anchor 범위 밖으로 끌어오지 않는다", test_merge_does_not_forward_fill)
    check("머리글 행을 찾는다", test_finds_header_row)
    check("LED 시트는 visible 이어도 이름으로 걸러진다", test_led_sheet_skipped_even_when_visible)
    check("노무비 수식 끝의 배율을 읽는다", test_reads_multiplier_from_formula)
    check("배율 없는 표준 수식에는 배율을 붙이지 않는다", test_no_multiplier_on_plain_formula)
    check("모르는 수식은 해석하지 않고 표시만 한다", test_unknown_formula_is_flagged_not_guessed)
    check("규격 문자열에서 배율을 유추하지 않는다", test_multiplier_is_not_guessed_from_spec)
    check("표준 금액 수식은 표시하지 않는다", test_plain_amounts_have_no_shape)
    check("시트 간 참조도 표준으로 본다", test_cross_sheet_amounts_are_still_plain)
    check("금액 칸이 전부 INT 면 int 로 표시한다", test_all_int_amounts_are_marked_int)
    check("일부만 INT 인 행은 mixed 다", test_mixed_amounts_are_not_int)
    check("상수로 덮인 칸은 constant 다", test_constant_amount_is_marked_constant)
    print()
    if _FAILURES:
        print(f"실패 {len(_FAILURES)} / 통과 {_PASSED}")
        for failure in _FAILURES:
            print(f"  - {failure}")
        raise SystemExit(1)
    print(f"통과 {_PASSED}")


if __name__ == "__main__":
    main()
