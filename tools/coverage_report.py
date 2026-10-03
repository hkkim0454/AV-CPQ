# -*- coding: utf-8 -*-
"""실제 견적서의 품목을 카탈로그로 얼마나 덮을 수 있는지 본다 (계획 Task 6).

카탈로그가 1,468개라도 **실제 견적에 쓰는 품목을 못 찾으면 쓸모가 없다.**
평택 견적서(설계서 §4.1 기준 파일)의 품목 행을 뽑아 `data/approved/products.json`과
대조해 적중률을 낸다.

금액은 읽지 않는다. 품명(B)·규격(C)·단위(D)만 본다 (설계서 §4.5).

사용:  python tools/coverage_report.py --source <견적서.xlsx> [출력.md]
"""
import json
import os
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parent.parent
PRODUCTS = ROOT / "data" / "approved" / "products.json"

# 견적서 내역 시트의 열 (docs/template/mapping.md §4.2)
COL_NAME, COL_SPEC, COL_UNIT = 2, 3, 4
FIRST_BODY_ROW = 5

# 품목이 아닌 행 — 소계·간접비·그룹 머리글
NON_ITEM = re.compile(
    r"^(직접비|간접비|합\s*계|소\s*계|직접비계|간접비계|Ⅰ|Ⅱ|\[.*\]|"
    r"간접노무비|고용보험료|산재보험료|연금보험료|건강보험료|노인장기요양보험료|"
    r"산업안전보건관리비|퇴직공제부금비|공과잡비)$"
)


def normalize(text):
    """대조용 정규화. 공백·대소문자·전각을 없앤다."""
    if text is None:
        return ""
    value = unicodedata.normalize("NFKC", str(text)).strip().lower()
    value = re.sub(r"^[-\s·]+", "", value)      # 딸림 품목의 앞 `- `
    value = re.sub(r"[\s_]+", "", value)
    return value


def load_catalog():
    data = json.loads(PRODUCTS.read_text(encoding="utf-8"))
    by_name = {}
    by_spec = {}
    for product in data["products"]:
        by_name.setdefault(normalize(product["quoteName"]), []).append(product)
        spec = normalize(product["quoteSpec"])
        if spec:
            by_spec.setdefault(spec, []).append(product)
    return data["products"], by_name, by_spec


def read_quote_items(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    items = []
    for ws in wb.worksheets:
        if ws.title == "갑지":
            continue
        for row in ws.iter_rows(min_row=FIRST_BODY_ROW, max_col=4, values_only=True):
            name = row[COL_NAME - 1]
            unit = row[COL_UNIT - 1]
            if name is None or unit is None:
                continue
            text = str(name).strip()
            if text == "" or NON_ITEM.match(text):
                continue
            items.append(
                {
                    "sheet": ws.title,
                    "name": text,
                    "spec": "" if row[COL_SPEC - 1] is None else str(row[COL_SPEC - 1]).strip(),
                    "unit": str(unit).strip(),
                }
            )
    wb.close()
    return items


def match(item, by_name, by_spec):
    """규격(모델명) 일치를 품명 일치보다 믿는다. 모델명이 더 특정적이다."""
    spec = normalize(item["spec"])
    if spec and spec in by_spec:
        return "spec", by_spec[spec][0]
    name = normalize(item["name"])
    if name in by_name:
        return "name", by_name[name][0]
    return None, None


MISS_REASONS = [
    ("파생 행", "다른 행의 금액에서 단가가 계산된다. 카탈로그 품목이 아니다"),
    ("지급자재", "고객이 주는 장비. 단가가 없고 카탈로그에 있을 이유가 없다"),
    ("LED 전광판", "숨김 `LED전광판 계산` 시트로 구성한다. 고르는 품목이 아니다"),
    ("현장 공사 항목", "설치비·프로그래밍비 등 프로젝트마다 다른 항목"),
    ("카탈로그 보충 후보", "실제로 카탈로그에 없는 제품. **이것만 보충 대상이다**"),
]

DERIVED = re.compile(r"잡자재비|배관\s*기타자재")
SUPPLIED = re.compile(r"지급자재")
LED_SHEET = re.compile(r"LED")
SITE_WORK = re.compile(r"설치비|설치인건비|programming\s*fee|세팅비|시험|조정|인테리어", re.I)


def classify_misses(misses):
    """미적중을 사람이 조치할 수 있는 분류로 나눈다."""
    buckets = {label: [] for label, _ in MISS_REASONS}
    for item in misses:
        name, spec, sheet = item["name"], item["spec"], item["sheet"]
        if DERIVED.search(name):
            buckets["파생 행"].append(item)
        elif SUPPLIED.search(spec) or SUPPLIED.search(name):
            buckets["지급자재"].append(item)
        elif LED_SHEET.search(sheet) and LED_SHEET.search(f"{name} {spec}"):
            buckets["LED 전광판"].append(item)
        elif SITE_WORK.search(f"{name} {spec}"):
            buckets["현장 공사 항목"].append(item)
        else:
            buckets["카탈로그 보충 후보"].append(item)
    return buckets


def main():
    args = sys.argv[1:]
    source = os.environ.get("AVCPQ_NEGO_XLSX")
    destination = None
    index = 0
    while index < len(args):
        if args[index] == "--source" and index + 1 < len(args):
            source = args[index + 1]
            index += 2
            continue
        destination = args[index]
        index += 1

    if not source:
        raise SystemExit("견적서 경로를 지정한다: --source <경로> 또는 AVCPQ_NEGO_XLSX")

    products, by_name, by_spec = load_catalog()
    items = read_quote_items(Path(source))

    hits_by_kind = Counter()
    misses = []
    for item in items:
        kind, product = match(item, by_name, by_spec)
        if product is None:
            misses.append(item)
        else:
            hits_by_kind[kind] += 1

    total = len(items)
    hit = sum(hits_by_kind.values())
    rate = (hit / total * 100) if total else 0.0

    lines = []
    add = lines.append
    add("# 카탈로그 커버리지 보고")
    add("")
    add("실제 견적서의 품목을 `data/approved/products.json`으로 얼마나 덮는지 측정했다.")
    add("계획 Task 6. 금액은 읽지 않았다 — 품명·규격·단위만 대조했다.")
    add("")
    add(f"- 카탈로그 제품: **{len(products)}**")
    add(f"- 견적서 품목 행: **{total}**")
    add(f"- 적중: **{hit} ({rate:.1f}%)**  — 규격 일치 {hits_by_kind['spec']}, 품명 일치 {hits_by_kind['name']}")
    add(f"- 미적중: **{len(misses)}**")
    add("")
    add("## 대조 방법")
    add("")
    add("규격(모델명) 일치를 품명 일치보다 우선한다. 모델명이 더 특정적이기 때문이다.")
    add("정규화는 NFKC + 소문자 + 공백·밑줄 제거 + 딸림 품목의 앞 `- ` 제거까지만 한다.")
    add("**부분 일치나 유사도 매칭을 쓰지 않는다** — 엉뚱한 제품에 원가·품셈이 붙으면")
    add("조용히 틀린 견적이 나온다 (설계서 §7.5와 같은 원칙).")
    add("")
    buckets = classify_misses(misses)
    real_gap = buckets["카탈로그 보충 후보"]
    adjusted_total = total - sum(
        len(v) for k, v in buckets.items() if k != "카탈로그 보충 후보"
    )
    adjusted_rate = (hit / adjusted_total * 100) if adjusted_total else 0.0

    add("## 미적중 분류")
    add("")
    add("적중률이 낮다고 카탈로그가 잘못된 것은 아니다. 카탈로그에 **있을 이유가 없는**")
    add("행이 섞여 있다. 사람이 보충해야 하는 것은 마지막 분류뿐이다.")
    add("")
    add("| 분류 | 건수 | 뜻 |")
    add("|---|---|---|")
    for label, reason in MISS_REASONS:
        add(f"| {label} | {len(buckets[label])} | {reason} |")
    add("")
    add(f"보충 대상만 분모로 잡으면 적중률은 **{adjusted_rate:.1f}%** ({hit}/{adjusted_total})다.")
    add("")

    for label, _ in MISS_REASONS:
        rows = buckets[label]
        if not rows:
            continue
        add(f"### {label} ({len(rows)})")
        add("")
        add("| 시트 | 품명 | 규격 | 단위 |")
        add("|---|---|---|---|")
        for item in rows:
            add(f"| {item['sheet']} | {item['name']} | {item['spec']} | {item['unit']} |")
        add("")

    add("## 확인한 사실")
    add("")
    add("- 숨김 시트 `LED전광판 계산`은 **제품 목록이 아니라 계산 시트**다.")
    add("  `000인치 삼성 LED전광판 (0x0)`, `LED 설치 (00.00㎡)`처럼 자리표시자로 되어 있고,")
    add("  실제 견적에서는 화면 크기·캐비닛 수를 계산해 그때그때 구성한다.")
    add("  → 카탈로그에서 제외한 것이 맞다. LED 전광판은 **고르는 것이 아니라 구성하는 것**이다.")
    add("- 숨김 시트 `단종, 미사용 제품`은 이름 그대로 단종품이다. 제외가 맞다.")
    add("")
    add("## 대조하지 않은 것")
    add("")
    add("가격·노무비는 대조하지 않았다. 카탈로그 단가가 실제 견적 단가와 같은지는")
    add("별도 확인이 필요하다 — 견적마다 조정이 들어가므로 다른 것이 정상일 수 있다.")

    if real_gap:
        print("\n카탈로그 보충 후보:")
        for item in real_gap:
            print(f"  - {item['name']}  ({item['spec']})")

    report = "\n".join(lines) + "\n"
    if destination:
        Path(destination).parent.mkdir(parents=True, exist_ok=True)
        Path(destination).write_text(report, encoding="utf-8")
        print(f"saved: {destination}")

    print(f"견적서 품목 {total} / 적중 {hit} ({rate:.1f}%) / 미적중 {len(misses)}")
    print(f"  규격 일치 {hits_by_kind['spec']}  품명 일치 {hits_by_kind['name']}")


if __name__ == "__main__":
    main()
